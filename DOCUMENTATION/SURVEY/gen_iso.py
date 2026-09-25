from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)

def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    return p

def add_question(doc, title, desc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(10.5)

    p2 = doc.add_paragraph()
    p2.paragraph_format.left_indent = Inches(0.2)
    p2.paragraph_format.space_after = Pt(3)
    r2 = p2.add_run(desc)
    r2.font.size = Pt(10)
    r2.italic = True

    tbl = doc.add_table(rows=2, cols=5)
    tbl.style = 'Table Grid'
    headers = ['5 - Strongly Agree', '4 - Agree', '3 - Neutral', '2 - Disagree', '1 - Strongly Disagree']
    for i, h in enumerate(headers):
        c = tbl.cell(0, i)
        c.text = h
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in c.paragraphs[0].runs:
            r.font.size = Pt(9)
            r.bold = True
        tbl.cell(1, i).text = ''
    doc.add_paragraph()

# ---- Title ----
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('ISO/IEC 25010:2023 SOFTWARE QUALITY EVALUATION INSTRUMENT')
r.bold = True
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run('Web-Based Data-Driven Job Recommendation System with Dashboard for PESO CSJDM')
r2.font.size = Pt(11)
r2.italic = True

doc.add_paragraph()

p3 = doc.add_paragraph()
r3 = p3.add_run('Instructions: ')
r3.bold = True
r3.font.size = Pt(10)
r4 = p3.add_run('Please assess each item using the following rating scale:')
r4.font.size = Pt(10)

p4 = doc.add_paragraph()
p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
r5 = p4.add_run('5 - Strongly Agree     4 - Agree     3 - Neutral     2 - Disagree     1 - Strongly Disagree')
r5.bold = True
r5.font.size = Pt(10)

doc.add_paragraph()

# ---- Respondent Info ----
info_fields = [
    'Capstone Project Title:',
    'Researchers/Proponents:',
    'Evaluator Name:',
    'Date:',
]
for f in info_fields:
    pi = doc.add_paragraph()
    ri = pi.add_run(f + '  ')
    ri.bold = True
    ri.font.size = Pt(10)
    pi.add_run('_' * 50)

doc.add_paragraph()

# ---- Sections ----
sections = [
    ('A. FUNCTIONAL SUITABILITY', [
        ('1. Functional Completeness',
         'The system provides all necessary functions to support PESO staff in managing applicant profiles, job vacancies, and generating job recommendations.'),
        ('2. Functional Correctness',
         'The system produces accurate job recommendations and analytics results based on applicant profiles and available job vacancies.'),
        ('3. Functional Appropriateness',
         "The system's features appropriately support PESO staff in matching applicants to suitable job opportunities."),
    ]),
    ('B. PERFORMANCE EFFICIENCY', [
        ('4. Time Behaviour',
         'The system responds promptly to user actions such as generating recommendations, searching records, and loading the analytics dashboard.'),
        ('5. Resource Utilization',
         'The system efficiently uses computing resources when processing applicant data and generating job recommendations.'),
        ('6. Capacity',
         'The system can handle the expected volume of applicant records and job vacancies without performance degradation.'),
    ]),
    ('C. COMPATIBILITY', [
        ('7. Co-existence',
         'The system operates effectively alongside other tools or software used in the PESO office environment.'),
        ('8. Interoperability',
         'The system can work with external data sources and formats used in employment service operations.'),
    ]),
    ('D. INTERACTION CAPABILITY', [
        ('9. Appropriateness Recognizability',
         'Users can easily understand the purpose of the system and its job recommendation features upon first use.'),
        ('10. Learnability',
         "PESO staff can quickly learn to navigate and use the system's features without extensive training."),
        ('11. Operability',
         "The system's functions for managing applicants, vacancies, and recommendations are easy to operate and control."),
        ('12. User Error Protection',
         'The system prevents or minimizes data entry errors when managing applicant profiles and job vacancies.'),
        ('13. User Engagement',
         "The system's interface is well-designed and encourages PESO staff to use it consistently in daily operations."),
        ('14. Inclusivity',
         'The system is usable by PESO staff regardless of their level of technical proficiency.'),
        ('15. User Assistance',
         'The system provides sufficient guidance to help users navigate and accomplish their tasks.'),
        ('16. Self-descriptiveness',
         "The system's interface and features are clear and intuitive, allowing users to understand how to use them without needing additional documentation."),
    ]),
    ('E. RELIABILITY', [
        ('17. Faultlessness',
         'The system consistently generates job recommendations and manages records without errors during normal operation.'),
        ('18. Availability',
         'The system is accessible and ready for use whenever PESO staff need it.'),
        ('19. Fault Tolerance',
         'The system continues to function normally even when encountering unexpected inputs or minor errors.'),
        ('20. Recoverability',
         'The system can recover applicant data and return to normal operation after an unexpected interruption or failure.'),
    ]),
    ('F. SECURITY', [
        ('21. Confidentiality',
         'The system ensures that applicant information and employment data are accessible only to authorized PESO staff.'),
        ('22. Integrity',
         'The system protects applicant records and recommendation data from unauthorized modification or deletion.'),
        ('23. Non-repudiation',
         'The system maintains records of significant actions performed, ensuring they can be verified and cannot be denied later.'),
        ('24. Accountability',
         'The system logs and tracks user actions, allowing activities to be traced back to the responsible staff member.'),
        ('25. Authenticity',
         'The system verifies the identity of users before granting access to applicant data and system features.'),
        ('26. Resistance',
         'The system remains secure and functional even when subjected to unauthorized access attempts.'),
    ]),
    ('G. MAINTAINABILITY', [
        ('27. Modularity',
         'The system is structured so that updates or changes to one feature do not disrupt other functions.'),
        ('28. Reusability',
         'Components or features of the system can be repurposed or extended to support future enhancements.'),
        ('29. Analysability',
         'Issues or defects in the system can be identified and diagnosed without difficulty.'),
        ('30. Modifiability',
         "The system can be updated or enhanced to accommodate changes in PESO's operational requirements without introducing new errors."),
        ('31. Testability',
         "The system's features can be systematically tested to verify they meet the defined requirements."),
    ]),
    ('H. FLEXIBILITY', [
        ('32. Adaptability',
         'The system can be deployed and operated in different environments or configurations as needed by PESO.'),
        ('33. Scalability',
         'The system can accommodate growth in the number of applicants, vacancies, or users without significant performance loss.'),
        ('34. Installability',
         "The system can be set up and deployed in PESO's operating environment with ease."),
        ('35. Replaceability',
         "The system's components can be updated or replaced to meet evolving needs without disrupting overall operations."),
    ]),
    ('I. SAFETY', [
        ('36. Operational Constraint',
         'The system operates within defined parameters to prevent unintended actions that could affect applicant data integrity.'),
        ('37. Risk Identification',
         'The system can identify and flag potentially problematic inputs or operations before they cause issues.'),
        ('38. Fail Safe',
         'The system defaults to a safe state when it encounters an error, preventing data loss or corruption.'),
        ('39. Hazard Warning',
         'The system alerts users to potential issues, such as incomplete applicant profiles or missing required data, before processing.'),
        ('40. Safe Integration',
         'The system maintains data integrity and safe operation when working with imported records or connected data sources.'),
    ]),
]

for sec_title, questions in sections:
    add_section_heading(doc, sec_title)
    for q_title, q_desc in questions:
        add_question(doc, q_title, q_desc)

# ---- Comments ----
doc.add_paragraph()
pc = doc.add_paragraph()
rc = pc.add_run('Comments / Suggestions (Optional):')
rc.bold = True
rc.font.size = Pt(10)
for _ in range(3):
    doc.add_paragraph('_' * 80)

doc.save(r'C:\Users\user\Documents\CAPSTONE\DOCUMENTATION\SURVEY\ISO_Evaluation_Form.docx')
print('ISO_Evaluation_Form.docx saved.')
