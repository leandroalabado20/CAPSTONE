from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)

def add_section_heading(doc, title, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Inches(0.1)
        p2.paragraph_format.space_after = Pt(4)
        r2 = p2.add_run(subtitle)
        r2.font.size = Pt(10)
        r2.italic = True

def add_question(doc, code, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(code + '  ')
    r.bold = True
    r.font.size = Pt(10.5)
    p.add_run(text).font.size = Pt(10.5)

    tbl = doc.add_table(rows=2, cols=5)
    tbl.style = 'Table Grid'
    headers = ['5 - Strongly Agree', '4 - Agree', '3 - Neutral', '2 - Disagree', '1 - Strongly Disagree']
    for i, h in enumerate(headers):
        c = tbl.cell(0, i)
        c.text = h
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in c.paragraphs[0].runs:
            r.font.size = Pt(9)
            r.bold = True
        tbl.cell(1, i).text = ''
    doc.add_paragraph()

# ---- Title ----
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('TECHNOLOGY ACCEPTANCE MODEL (TAM) EVALUATION INSTRUMENT')
r.bold = True
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run('Web-Based Data-Driven Job Recommendation System with Dashboard for PESO CSJDM')
r2.font.size = Pt(11)
r2.italic = True

doc.add_paragraph()

intro = doc.add_paragraph()
ri = intro.add_run(
    'Dear Respondent, please evaluate the system you just used by rating each statement below '
    'based on your honest experience and observation. Your responses will be treated with strict '
    'confidentiality and used solely for academic research purposes.'
)
ri.font.size = Pt(10)
intro.paragraph_format.space_after = Pt(6)

p3 = doc.add_paragraph()
r3 = p3.add_run('Rating Scale: ')
r3.bold = True
r3.font.size = Pt(10)
r4 = p3.add_run('5 - Strongly Agree     4 - Agree     3 - Neutral     2 - Disagree     1 - Strongly Disagree')
r4.font.size = Pt(10)

doc.add_paragraph()

# ---- Respondent Info ----
add_section_heading(doc, 'Respondent Profile')

info_fields = [
    'Capstone Project Title:',
    'Researchers/Proponents:',
    'Age:',
    'Sex:',
    'Year Level / Position:',
    'Frequency of System Use:',
]
for f in info_fields:
    pi = doc.add_paragraph()
    ri = pi.add_run(f + '  ')
    ri.bold = True
    ri.font.size = Pt(10)
    pi.add_run('_' * 45)

doc.add_paragraph()

# ---- A. Perceived Usefulness ----
add_section_heading(
    doc,
    'A. Perceived Usefulness (PU)',
    'This measures whether the user believes the system helps them accomplish tasks effectively.'
)
pu_items = [
    ('PU1.', 'Using this system improves my performance in accomplishing job referral and applicant management tasks.'),
    ('PU2.', 'Using this system increases my productivity when working with job applicants and vacancies.'),
    ('PU3.', 'This system helps me identify suitable job category matches for applicants more quickly.'),
    ('PU4.', 'This system is useful for generating job category recommendations for applicants.'),
    ('PU5.', 'Overall, I find this system beneficial for job referral and employment matching purposes.'),
]
for code, text in pu_items:
    add_question(doc, code, text)

# ---- B. Perceived Ease of Use ----
add_section_heading(
    doc,
    'B. Perceived Ease of Use (PEOU)',
    'This measures whether the user finds the system simple and easy to operate.'
)
peou_items = [
    ('PEOU1.', "Learning to use this system's features, such as registering applicants and generating job recommendations, was easy for me."),
    ('PEOU2.', 'I find it easy to navigate the system and perform tasks such as managing applicant records or viewing job recommendations.'),
    ('PEOU3.', 'My interaction with the system, including its forms, navigation, and dashboard, is clear and understandable.'),
    ('PEOU4.', 'I find the system easy to use overall.'),
    ('PEOU5.', "The system's interface (buttons, menus, navigation) is user-friendly."),
]
for code, text in peou_items:
    add_question(doc, code, text)

# ---- C. Behavioral Intention to Use ----
add_section_heading(
    doc,
    'C. Behavioral Intention to Use (BI)',
    'This measures whether the user intends to continue using the system going forward.'
)
bi_items = [
    ('BI1.', 'I intend to use this system regularly if it is made available.'),
    ('BI2.', 'I would recommend this system to others who work with job applicants and employment referrals.'),
    ('BI3.', 'I plan to use this system regularly for applicant management and job matching if it becomes available.'),
    ('BI4.', 'Given the chance, I would prefer using this system over manual methods for managing applicants and generating job recommendations.'),
]
for code, text in bi_items:
    add_question(doc, code, text)

# ---- D. Open-Ended ----
add_section_heading(doc, 'D. Open-Ended Questions (Optional)')
open_items = [
    '1. What did you like most about the system?',
    '2. What improvements would you suggest?',
]
for q in open_items:
    po = doc.add_paragraph()
    ro = po.add_run(q)
    ro.bold = True
    ro.font.size = Pt(10.5)
    for _ in range(3):
        doc.add_paragraph('_' * 80)
    doc.add_paragraph()

doc.save(r'C:\Users\user\Documents\CAPSTONE\DOCUMENTATION\SURVEY\TAM_Evaluation_Form.docx')
print('TAM_Evaluation_Form.docx saved.')
